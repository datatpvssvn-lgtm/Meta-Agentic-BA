from __future__ import annotations

from pathlib import Path
import argparse
import json
import sys

import openpyxl

sys.stdout.reconfigure(encoding="utf-8")


ROOT = Path(__file__).resolve().parents[3]


def main() -> None:
    parser = argparse.ArgumentParser(description="Inspect selected Joycat sources from a portable root.")
    parser.add_argument("--root", type=Path, default=ROOT)
    args = parser.parse_args()
    root = args.root.expanduser().resolve()
    files = [
        root / "01_inputs/joycat/JOYCAT_CAMPAIGN_3_THANG_CO_CHI_PHI_OBJECTIVE_DEMO.xlsx",
        root / "01_inputs/joycat/JOYCAT_SHOPEE_PRODUCTS_2026-08-25.xlsx",
        root / "01_inputs/joycat/raw/meta_ads/preferred_candidate/CPAS-SHOPEE-Campaigns-1-Mar-2026-31-Mar-2026.xlsx",
        root / "01_inputs/joycat/raw/meta_ads/preferred_candidate/CPAS-SHOPEE-Campaigns-1-Apr-2026-30-Apr-2026.xlsx",
    ]
    for path in files:
        workbook = openpyxl.load_workbook(path, data_only=False, read_only=False)
        print(json.dumps({"file": str(path), "sheets": workbook.sheetnames}, ensure_ascii=False))
        for sheet in workbook.worksheets:
            print(f"SHEET\t{sheet.title}\trows={sheet.max_row}\tcols={sheet.max_column}")
            for row in sheet.iter_rows(min_row=1, max_row=min(sheet.max_row, 8), values_only=True):
                print(repr(row))
            if "Apr-2026" in path.name:
                for row_number in range(1, sheet.max_row + 1):
                    name = sheet.cell(row_number, 3).value
                    if name and "9/4/2026" in str(name):
                        print("APR_MATCH", row_number, repr(name), sheet.cell(row_number, 10).value)


if __name__ == "__main__":
    main()
