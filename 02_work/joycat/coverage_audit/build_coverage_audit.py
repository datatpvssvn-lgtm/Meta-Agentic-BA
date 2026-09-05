from __future__ import annotations

from collections import defaultdict
from pathlib import Path
import csv
import json
import re
import sys

import openpyxl


sys.stdout.reconfigure(encoding="utf-8")

ROOT = Path(r"D:\Meta Agentic BA")
WORK = ROOT / "02_work" / "joycat" / "coverage_audit"
PREFERRED = ROOT / "01_inputs" / "joycat" / "raw" / "meta_ads" / "preferred_candidate"
DEMO = ROOT / "03_outputs" / "joycat" / "JOYCAT_CAMPAIGN_3_THANG_CO_CHI_PHI_OBJECTIVE_DEMO.xlsx"
OUT_CSV = WORK / "coverage_6_pairs_detail.csv"
OUT_JSON = WORK / "coverage_audit_summary.json"

MONTHS = ["2026-03", "2026-04", "2026-05"]
CAMPAIGN_FILES = {
    "2026-03": "CPAS-SHOPEE-Campaigns-1-Mar-2026-31-Mar-2026.xlsx",
    "2026-04": "CPAS-SHOPEE-Campaigns-1-Apr-2026-30-Apr-2026.xlsx",
    "2026-05": "CPAS-SHOPEE-Campaigns-1-May-2026-31-May-2026.xlsx",
}
OBJECTIVES = [
    "Engagement - Messaging",
    "Engagement - Post engagement",
    "Sales - Purchase",
    "Sales - Purchase [Suy luận từ tên Campaign]",
    "Awareness - Ad recall",
    "Chưa xác định",
]
FUNNELS = ["TOFU", "MOFU", "BOFU"]
PRODUCT_BUCKETS = [
    ("Nhóm sản phẩm", "Vi sinh khử mùi"),
    ("Nhóm sản phẩm", "Cát Sắn Natural"),
    ("Nhóm sản phẩm", "Cát Mix PerformX"),
    ("Nhóm sản phẩm", "Cát Khoáng MineralX"),
    ("Nhóm sản phẩm", "Cát Đậu Nành Mix Khoáng PowerX"),
    ("Nhóm sản phẩm", "Cát Sắn Mix Đậu Nành Tofuu.X"),
    ("Ngành hàng", "Cát nói chung"),
    ("Shared/Unmapped", "Shared/Unmapped"),
]


def normalize_header(value: object) -> str:
    return "" if value is None else str(value).strip()


def sheet_rows(path: Path, sheet_name: str | None = None, header_row: int = 1):
    workbook = openpyxl.load_workbook(path, data_only=True, read_only=False)
    sheet = workbook[sheet_name] if sheet_name else workbook[workbook.sheetnames[0]]
    headers = [normalize_header(cell.value) for cell in sheet[header_row]]
    records = []
    for row_number in range(header_row + 1, sheet.max_row + 1):
        values = [sheet.cell(row_number, col).value for col in range(1, len(headers) + 1)]
        if not any(value is not None and str(value).strip() != "" for value in values):
            continue
        records.append((row_number, dict(zip(headers, values))))
    return sheet.title, records


def product_bucket(campaign_name: str) -> tuple[str, str, str]:
    name = (campaign_name or "").upper()
    if "VI SINH" in name:
        return "Nhóm sản phẩm", "Vi sinh khử mùi", "Tên Campaign chứa VI SINH"
    if re.search(r"(^|[^A-Z0-9])PFX([^A-Z0-9]|$)", name):
        return "Nhóm sản phẩm", "Cát Mix PerformX", "Tên Campaign chứa PFX"
    if re.search(r"(^|[^A-Z0-9])MNX([^A-Z0-9]|$)", name):
        return "Nhóm sản phẩm", "Cát Khoáng MineralX", "Tên Campaign chứa MNX"
    if re.search(r"(^|[^A-Z0-9])NA([^A-Z0-9]|$)", name):
        return "Nhóm sản phẩm", "Cát Sắn Natural", "Tên Campaign chứa NA"
    if "CÁT" in name:
        return "Ngành hàng", "Cát nói chung", "Tên Campaign chỉ nói CÁT; không đủ chọn nhóm/listing"
    return "Shared/Unmapped", "Shared/Unmapped", "Tên Campaign không đủ căn cứ gắn sản phẩm"


demo_sheet, demo_records = sheet_rows(DEMO, "Campaign_3_tháng", 5)
rows = []
for source_row, record in demo_records:
    spend = record.get("Amount spent (VND)")
    if not isinstance(spend, (int, float)) or spend <= 0:
        continue
    product_level, product_value, product_basis = product_bucket(str(record.get("Campaign name") or ""))
    rows.append({
        "period": str(record.get("Tháng")),
        "campaign_name": str(record.get("Campaign name") or ""),
        "result_indicator": record.get("Result indicator"),
        "cost_per_result": record.get("Cost per results"),
        "spend": float(spend),
        "funnel": str(record.get("Phễu") or "Chưa phân loại"),
        "objective": str(record.get("Objective suy luận") or "Chưa xác định"),
        "product_level": product_level,
        "product_value": product_value,
        "product_basis": product_basis,
        "demo_row": source_row,
    })


raw_control = {}
source_lookup = {}
source_names_by_period = defaultdict(list)
for period, filename in CAMPAIGN_FILES.items():
    path = PREFERRED / filename
    sheet_name, records = sheet_rows(path, None, 1)
    total_row_spend = None
    positive_rows = []
    for row_number, record in records:
        spend = record.get("Amount spent (VND)")
        name = record.get("Campaign name")
        if name in (None, "") and isinstance(spend, (int, float)):
            total_row_spend = float(spend)
        if name not in (None, "") and isinstance(spend, (int, float)) and spend > 0:
            positive_rows.append((row_number, record))
            source_lookup[(period, str(name))] = {
                "file": filename,
                "sheet": sheet_name,
                "row": row_number,
                "record": record,
            }
            source_names_by_period[period].append(str(name))
    raw_control[period] = {
        "file": filename,
        "sheet": sheet_name,
        "positive_campaign_rows": len(positive_rows),
        "sum_positive_campaign_spend": sum(float(record["Amount spent (VND)"]) for _, record in positive_rows),
        "reported_total_row_spend": total_row_spend,
        "positive_spend_blank_result_rows": sum(
            1 for _, record in positive_rows
            if record.get("Results") in (None, "") and record.get("Result indicator") in (None, "")
        ),
    }


for row in rows:
    match = source_lookup.get((row["period"], row["campaign_name"]))
    row["raw_source"] = match


def aggregate(key_fields: tuple[str, ...]):
    grouped = defaultdict(lambda: {"row_count": 0, "spend": 0.0, "demo_rows": [], "campaigns": []})
    for row in rows:
        key = tuple(row[field] for field in key_fields)
        grouped[key]["row_count"] += 1
        grouped[key]["spend"] += row["spend"]
        grouped[key]["demo_rows"].append(row["demo_row"])
        grouped[key]["campaigns"].append(row["campaign_name"])
    return grouped


funnel_product = aggregate(("period", "funnel", "product_level", "product_value"))
product_objective = aggregate(("period", "product_level", "product_value", "objective"))
funnel_objective = aggregate(("period", "funnel", "objective"))


coverage_rows = []


def add_coverage(
    period: str,
    pair: str,
    dim_1: str,
    value_1: str,
    dim_2: str,
    value_2: str,
    product_level: str = "",
    group: dict | None = None,
    status: str = "",
    reason: str = "",
    next_action: str = "",
    mapping_version: str = "",
):
    coverage_rows.append({
        "Kỳ dữ liệu": period,
        "Cặp phân tích": pair,
        "Chiều 1": dim_1,
        "Giá trị chiều 1": value_1,
        "Chiều 2": dim_2,
        "Giá trị chiều 2": value_2,
        "Cấp sản phẩm": product_level,
        "Số dòng nguồn": "" if group is None else group["row_count"],
        "Ads Cost (VND)": "" if group is None else round(group["spend"]),
        "Nguồn": "JOYCAT_CAMPAIGN_3_THANG_CO_CHI_PHI_OBJECTIVE_DEMO.xlsx",
        "Sheet": demo_sheet,
        "Dòng nguồn": "" if group is None else ",".join(str(x) for x in group["demo_rows"]),
        "Mapping version": mapping_version,
        "Trạng thái coverage": status,
        "Lý do": reason,
        "Bước xử lý": next_action,
    })


for period in MONTHS:
    for pair, dim_2 in [
        ("Nền tảng × Sản phẩm", "Sản phẩm"),
        ("Nền tảng × Phễu", "Phễu"),
        ("Nền tảng × Campaign objective", "Campaign objective"),
    ]:
        add_coverage(
            period, pair, "Publisher platform", "Không có trong source", dim_2, "Chưa kiểm tra được",
            status="Thiếu nguồn/thiếu chiều",
            reason="Dataset hiện tại không có Publisher platform; không được tự gán Facebook/Instagram/WhatsApp.",
            next_action="Xin breakdown Publisher platform cùng kỳ và khóa nối/grain tương thích.",
        )

for period in MONTHS:
    for funnel in FUNNELS:
        for product_level, product_value in PRODUCT_BUCKETS:
            group = funnel_product.get((period, funnel, product_level, product_value))
            if group:
                status = "Tính được một phần bằng human mapping"
                reason = "Phễu và nhóm sản phẩm được gắn/suy luận trên cùng dòng Campaign; chưa xác minh tới listing/SKU."
                next_action = "Khóa rule/version Phễu và owner duyệt mapping nhóm sản phẩm; xin Ads ID → Item ID/SKU nếu cần xuống listing."
            else:
                status = "Chưa quan sát tổ hợp trong phạm vi đã kiểm tra"
                reason = "Không có dòng Campaign được gắn vào tổ hợp này; không được đổi thành Ads Cost = 0."
                next_action = "Giữ giá trị trống; xác nhận lại sau khi mapping được duyệt hoặc có nguồn chi tiết hơn."
            add_coverage(
                period, "Phễu × Sản phẩm", "Phễu", funnel, "Sản phẩm", product_value,
                product_level=product_level, group=group, status=status, reason=reason, next_action=next_action,
                mapping_version="Phễu: chưa khóa; Sản phẩm: inference_name_v0",
            )

for period in MONTHS:
    for product_level, product_value in PRODUCT_BUCKETS:
        for objective in OBJECTIVES:
            group = product_objective.get((period, product_level, product_value, objective))
            if group:
                status = "Tính được một phần bằng human mapping"
                reason = "Objective suy luận được Duy cho phép dùng; sản phẩm mới suy luận tới nhóm/ngành hàng hoặc Shared."
                next_action = "Giữ metadata Objective; owner duyệt product mapping hoặc xin Ads ID → Item ID/SKU."
            else:
                status = "Chưa quan sát tổ hợp trong phạm vi đã kiểm tra"
                reason = "Không có dòng Campaign được gắn vào tổ hợp này; không được đổi thành Ads Cost = 0."
                next_action = "Giữ giá trị trống; kiểm tra lại khi mapping/source thay đổi."
            add_coverage(
                period, "Sản phẩm × Campaign objective", "Sản phẩm", product_value,
                "Campaign objective", objective, product_level=product_level, group=group,
                status=status, reason=reason, next_action=next_action,
                mapping_version="Objective: human-curated hiện hành; Sản phẩm: inference_name_v0",
            )

for period in MONTHS:
    for funnel in FUNNELS:
        for objective in OBJECTIVES:
            group = funnel_objective.get((period, funnel, objective))
            if group:
                status = "Tính được bằng human mapping"
                reason = "Phễu và Objective suy luận nằm trên cùng dòng Campaign có spend."
                next_action = "Khóa rule/version/owner Phễu; giữ basis/version/reviewer của Objective."
            else:
                status = "Chưa quan sát tổ hợp trong phạm vi đã kiểm tra"
                reason = "Không có dòng Campaign được gắn vào tổ hợp này; không được đổi thành Ads Cost = 0."
                next_action = "Giữ giá trị trống; kiểm tra lại khi mapping/source thay đổi."
            add_coverage(
                period, "Phễu × Campaign objective", "Phễu", funnel,
                "Campaign objective", objective, group=group, status=status, reason=reason,
                next_action=next_action,
                mapping_version="Objective: human-curated hiện hành; Phễu: chưa khóa",
            )


WORK.mkdir(parents=True, exist_ok=True)
with OUT_CSV.open("w", encoding="utf-8-sig", newline="") as handle:
    writer = csv.DictWriter(handle, fieldnames=list(coverage_rows[0].keys()))
    writer.writeheader()
    writer.writerows(coverage_rows)


demo_control = {}
for period in MONTHS:
    period_rows = [row for row in rows if row["period"] == period]
    demo_control[period] = {
        "rows": len(period_rows),
        "spend": sum(row["spend"] for row in period_rows),
        "raw_matches": sum(1 for row in period_rows if row["raw_source"] is not None),
        "raw_unmatched": [row["campaign_name"] for row in period_rows if row["raw_source"] is None],
        "raw_unmatched_near_names": {
            row["campaign_name"]: [
                name for name in source_names_by_period[period]
                if row["campaign_name"].rstrip(" -") == name.rstrip(" -")
            ]
            for row in period_rows if row["raw_source"] is None
        },
        "raw_unmatched_details": [
            {
                "demo_row": row["demo_row"],
                "campaign_name": row["campaign_name"],
                "spend": row["spend"],
                "funnel": row["funnel"],
                "objective": row["objective"],
            }
            for row in period_rows if row["raw_source"] is None
        ],
        "preferred_campaigns_not_in_demo": [
            name for name in source_names_by_period[period]
            if name not in {row["campaign_name"] for row in period_rows}
        ],
        "matched_spend_differences": [
            {
                "demo_row": row["demo_row"],
                "campaign_name": row["campaign_name"],
                "demo_spend": row["spend"],
                "preferred_row": row["raw_source"]["row"],
                "preferred_spend": row["raw_source"]["record"].get("Amount spent (VND)"),
                "difference": row["spend"] - float(row["raw_source"]["record"].get("Amount spent (VND)")),
            }
            for row in period_rows
            if row["raw_source"] is not None
            and row["spend"] != float(row["raw_source"]["record"].get("Amount spent (VND)"))
        ],
    }


pair_reconciliation = {}
for pair in ["Phễu × Sản phẩm", "Sản phẩm × Campaign objective", "Phễu × Campaign objective"]:
    pair_reconciliation[pair] = {}
    for period in MONTHS:
        observed = [
            row for row in coverage_rows
            if row["Cặp phân tích"] == pair and row["Kỳ dữ liệu"] == period and row["Ads Cost (VND)"] != ""
        ]
        pair_reconciliation[pair][period] = {
            "observed_cells": len(observed),
            "mapped_rows": sum(int(row["Số dòng nguồn"]) for row in observed),
            "mapped_spend": sum(float(row["Ads Cost (VND)"]) for row in observed),
            "demo_rows": demo_control[period]["rows"],
            "demo_spend": demo_control[period]["spend"],
        }


example_name = "10/11/2025 - MOF - TESTING A/B - 1% - CÁT"
example = next(row for row in rows if row["period"] == "2026-03" and row["campaign_name"] == example_name)
raw = example["raw_source"]
example_payload = {
    "demo_file": DEMO.name,
    "demo_sheet": demo_sheet,
    "demo_row": example["demo_row"],
    "raw_file": raw["file"] if raw else None,
    "raw_sheet": raw["sheet"] if raw else None,
    "raw_row": raw["row"] if raw else None,
    "period": example["period"],
    "campaign_name": example["campaign_name"],
    "result_indicator": example["result_indicator"],
    "results": None if not raw else raw["record"].get("Results"),
    "cost_per_result": example["cost_per_result"],
    "amount_spent_vnd": example["spend"],
    "impressions": None if not raw else raw["record"].get("Impressions"),
    "reach": None if not raw else raw["record"].get("Reach"),
    "attribution_setting": None if not raw else raw["record"].get("Attribution setting"),
    "funnel": example["funnel"],
    "objective_inferred": example["objective"],
    "objective_basis": "Result indicator = actions:onsite_conversion.messaging_conversation_started_7d",
    "product_level": example["product_level"],
    "product_value": example["product_value"],
    "product_basis": example["product_basis"],
}

comparison_name = "10/11/2025 - MOF - TESTING A/B - GỐC - CÁT"
comparison = next(row for row in rows if row["period"] == "2026-03" and row["campaign_name"] == comparison_name)
comparison_raw = comparison["raw_source"]
comparison_payload = {
    "demo_row": comparison["demo_row"],
    "raw_file": comparison_raw["file"] if comparison_raw else None,
    "raw_sheet": comparison_raw["sheet"] if comparison_raw else None,
    "raw_row": comparison_raw["row"] if comparison_raw else None,
    "campaign_name": comparison["campaign_name"],
    "result_indicator": comparison["result_indicator"],
    "results": None if not comparison_raw else comparison_raw["record"].get("Results"),
    "cost_per_result": comparison["cost_per_result"],
    "amount_spent_vnd": comparison["spend"],
    "impressions": None if not comparison_raw else comparison_raw["record"].get("Impressions"),
    "reach": None if not comparison_raw else comparison_raw["record"].get("Reach"),
    "attribution_setting": None if not comparison_raw else comparison_raw["record"].get("Attribution setting"),
    "funnel": comparison["funnel"],
    "objective_inferred": comparison["objective"],
    "product_level": comparison["product_level"],
    "product_value": comparison["product_value"],
}


summary = {
    "generated_from": {
        "demo": str(DEMO.relative_to(ROOT)),
        "preferred_campaign_folder": str(PREFERRED.relative_to(ROOT)),
    },
    "demo_control": demo_control,
    "raw_control": raw_control,
    "pair_reconciliation": pair_reconciliation,
    "coverage_row_count": len(coverage_rows),
    "observed_values": {
        "funnels": sorted({row["funnel"] for row in rows}),
        "objectives": sorted({row["objective"] for row in rows}),
        "objective_row_counts": {
            value: sum(1 for row in rows if row["objective"] == value)
            for value in sorted({row["objective"] for row in rows})
        },
        "objective_spend": {
            value: sum(row["spend"] for row in rows if row["objective"] == value)
            for value in sorted({row["objective"] for row in rows})
        },
    },
    "example": example_payload,
    "comparison_candidate": comparison_payload,
}
OUT_JSON.write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8")
print(json.dumps(summary, ensure_ascii=False, indent=2))
