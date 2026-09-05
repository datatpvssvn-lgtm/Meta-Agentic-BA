from __future__ import annotations

import json
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


RAW_ROOT = Path(r"D:\Meta Agentic BA\01_inputs\joycat\raw")
OUTPUT = Path(r"D:\Meta Agentic BA\02_work\joycat\dataset_audit\schema_inventory.json")


def clean(value: Any) -> str | None:
    if value is None:
        return None
    text = str(value).strip()
    return text if text else None


def choose_header(rows: list[list[Any]]) -> tuple[int | None, list[str]]:
    candidates: list[tuple[float, int, list[str]]] = []
    for row_number, row in enumerate(rows, start=1):
        values = [clean(v) for v in row]
        non_empty = [v for v in values if v is not None]
        if len(non_empty) < 2:
            continue
        text_like = sum(1 for v in non_empty if not v.replace(".", "", 1).isdigit())
        unique_ratio = len(set(non_empty)) / len(non_empty)
        score = len(non_empty) + (text_like / len(non_empty)) * 2 + unique_ratio
        candidates.append((score, row_number, values))
    if not candidates:
        return None, []
    _, row_number, values = max(candidates, key=lambda item: (item[0], -item[1]))
    last = max((i for i, v in enumerate(values) if v is not None), default=-1)
    headers = [v if v is not None else f"__blank_col_{i + 1}" for i, v in enumerate(values[: last + 1])]
    return row_number, headers


def inspect_workbook(path: Path) -> dict[str, Any]:
    workbook = load_workbook(path, read_only=False, data_only=False)
    sheet_records: list[dict[str, Any]] = []
    for sheet in workbook.worksheets:
        preview_rows = [list(row) for row in sheet.iter_rows(min_row=1, max_row=min(sheet.max_row, 25), values_only=True)]
        header_row, headers = choose_header(preview_rows)
        non_empty_counts: dict[str, int] = {}
        if header_row and headers:
            counts = [0] * len(headers)
            for row in sheet.iter_rows(min_row=header_row + 1, values_only=True):
                for index, value in enumerate(row[: len(headers)]):
                    if clean(value) is not None:
                        counts[index] += 1
            non_empty_counts = {headers[i]: counts[i] for i in range(len(headers))}

        top_rows = []
        for idx, row in enumerate(preview_rows[:8], start=1):
            values = [clean(v) for v in row]
            last = max((i for i, v in enumerate(values) if v is not None), default=-1)
            if last >= 0:
                top_rows.append({"row": idx, "values": values[: last + 1]})

        sheet_records.append(
            {
                "sheet": sheet.title,
                "max_row": sheet.max_row,
                "max_column": sheet.max_column,
                "header_row": header_row,
                "headers": headers,
                "non_empty_counts": non_empty_counts,
                "top_rows": top_rows,
            }
        )
    workbook.close()
    return {
        "path": str(path),
        "relative_path": str(path.relative_to(RAW_ROOT)),
        "size_bytes": path.stat().st_size,
        "sheets": sheet_records,
    }


def main() -> None:
    files = sorted(
        [
            path
            for path in RAW_ROOT.rglob("*")
            if path.is_file() and path.suffix.lower() in {".xlsx", ".xlsm", ".xls", ".csv", ".tsv"}
        ],
        key=lambda path: str(path).lower(),
    )
    inventory: dict[str, Any] = {"raw_root": str(RAW_ROOT), "file_count": len(files), "files": []}
    for path in files:
        if path.suffix.lower() in {".xlsx", ".xlsm"}:
            try:
                inventory["files"].append(inspect_workbook(path))
            except Exception as exc:
                inventory["files"].append(
                    {"path": str(path), "relative_path": str(path.relative_to(RAW_ROOT)), "error": repr(exc)}
                )
        else:
            inventory["files"].append(
                {
                    "path": str(path),
                    "relative_path": str(path.relative_to(RAW_ROOT)),
                    "error": "Định dạng chưa được đọc trong script này",
                }
            )
    OUTPUT.write_text(json.dumps(inventory, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps({"file_count": len(files), "output": str(OUTPUT)}, ensure_ascii=False))


if __name__ == "__main__":
    main()
