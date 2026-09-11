from __future__ import annotations

import json
import re
import argparse
from collections import Counter, defaultdict
from datetime import date, datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


WORKSPACE = Path(__file__).resolve().parents[3]
RAW_ROOT = WORKSPACE / "01_inputs" / "joycat" / "raw"
OUTPUT_ROOT = WORKSPACE / "01_inputs" / "joycat"
WORK_ROOT = WORKSPACE / "02_work" / "joycat" / "dataset_audit"
OUTPUT_JSON = WORK_ROOT / "source_audit_v2.json"

SELECTED_OUTPUTS = [
    OUTPUT_ROOT / "JOYCAT_CAMPAIGN_3_THANG_CO_CHI_PHI.xlsx",
    OUTPUT_ROOT / "JOYCAT_CAMPAIGN_3_THANG_CO_CHI_PHI_OBJECTIVE_DEMO.xlsx",
    OUTPUT_ROOT / "JOYCAT_SHOPEE_PRODUCTS_2026-08-25.xlsx",
]

ID_FIELDS = {"account id", "campaign id", "ad set id", "ad id", "item id", "sku", "product id"}
TARGET_FIELDS = {
    "campaign objective",
    "objective",
    "objective suy luận",
    "optimization goal",
    "performance goal",
    "publisher platform",
    "platform position",
    "placement",
    "destination",
    "conversion location",
    "phễu",
    "funnel",
    "product",
    "product name",
    "sku",
    "item id",
    "campaign id",
    "ad set id",
    "ad id",
}


def text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    return str(value).strip()


def norm(value: Any) -> str:
    return re.sub(r"\s+", " ", text(value)).casefold()


def choose_header(ws) -> tuple[int | None, list[str]]:
    best: tuple[float, int, list[str]] | None = None
    for row_number in range(1, min(ws.max_row, 30) + 1):
        values = [text(ws.cell(row_number, col).value) for col in range(1, ws.max_column + 1)]
        non_empty = [v for v in values if v]
        if len(non_empty) < 2:
            continue
        keyword_hits = sum(
            1
            for v in non_empty
            if any(
                token in v.casefold()
                for token in (
                    "name",
                    "id",
                    "results",
                    "spent",
                    "reporting",
                    "objective",
                    "phễu",
                    "product",
                    "item",
                )
            )
        )
        unique_ratio = len(set(non_empty)) / len(non_empty)
        score = len(non_empty) + keyword_hits * 4 + unique_ratio
        candidate = (score, row_number, values)
        if best is None or candidate[0] > best[0] or (candidate[0] == best[0] and row_number < best[1]):
            best = candidate
    if best is None:
        return None, []
    _, row_number, values = best
    last = max(i for i, value in enumerate(values) if value)
    headers = [value or f"__blank_col_{index + 1}" for index, value in enumerate(values[: last + 1])]
    return row_number, headers


def canonical_field(headers: list[str], candidates: list[str]) -> int | None:
    wanted = {candidate.casefold() for candidate in candidates}
    for index, header in enumerate(headers, start=1):
        if header.casefold() in wanted:
            return index
    return None


def month_from_path(path: Path) -> str | None:
    name = path.name.casefold()
    if "mar-2026" in name or "032026" in name:
        return "2026-03"
    if "apr-2026" in name or "042026" in name:
        return "2026-04"
    if "may-2026" in name:
        return "2026-05"
    return None


def grain_from_path(path: Path) -> str | None:
    name = path.name.casefold()
    if "campaigns" in name:
        return "Campaign"
    if "ad-sets" in name or "adset" in str(path).casefold():
        return "Ad set"
    if "ads-" in name or "ads\\" in str(path).casefold():
        return "Ad"
    return None


def inspect_id_column(ws, header_row: int, col: int, header: str) -> dict[str, Any]:
    type_counts: Counter[str] = Counter()
    values: list[str] = []
    numeric_count = 0
    text_count = 0
    trailing_zero_numeric = 0
    scientific_display_count = 0
    for row in range(header_row + 1, ws.max_row + 1):
        cell = ws.cell(row, col)
        value = cell.value
        if value is None or text(value) == "":
            continue
        type_counts[cell.data_type] += 1
        rendered = text(value)
        if len(values) < 8:
            values.append(rendered)
        if cell.data_type == "n" or isinstance(value, (int, float)):
            numeric_count += 1
            if rendered.endswith("000") or rendered.endswith("0000"):
                trailing_zero_numeric += 1
        else:
            text_count += 1
        if "E+" in rendered.upper():
            scientific_display_count += 1
    return {
        "header": header,
        "non_empty": sum(type_counts.values()),
        "cell_data_types": dict(type_counts),
        "numeric_count": numeric_count,
        "text_count": text_count,
        "numeric_values_ending_000_or_0000": trailing_zero_numeric,
        "scientific_values": scientific_display_count,
        "samples": values,
        "risk": (
            "ID stored as numeric; exact 17–18 digit value may already be lost"
            if numeric_count
            else "ID stored as text in inspected cells"
        ),
    }


def inspect_sheet(path: Path, ws) -> dict[str, Any]:
    header_row, headers = choose_header(ws)
    record: dict[str, Any] = {
        "sheet": ws.title,
        "max_row": ws.max_row,
        "max_column": ws.max_column,
        "header_row": header_row,
        "headers": headers,
    }
    if header_row is None:
        return record

    non_empty = [0] * len(headers)
    data_types: list[Counter[str]] = [Counter() for _ in headers]
    distinct: list[set[str]] = [set() for _ in headers]
    sample_values: list[list[str]] = [[] for _ in headers]
    rows_with_any_data = 0

    for row in range(header_row + 1, ws.max_row + 1):
        row_has_data = False
        for index in range(len(headers)):
            cell = ws.cell(row, index + 1)
            value = cell.value
            value_text = text(value)
            if not value_text:
                continue
            row_has_data = True
            non_empty[index] += 1
            data_types[index][cell.data_type] += 1
            distinct[index].add(value_text)
            if len(sample_values[index]) < 5 and value_text not in sample_values[index]:
                sample_values[index].append(value_text)
        rows_with_any_data += int(row_has_data)

    record["data_rows_with_any_value"] = rows_with_any_data
    record["columns"] = [
        {
            "header": headers[index],
            "non_empty": non_empty[index],
            "blank": max(rows_with_any_data - non_empty[index], 0),
            "blank_rate": round(max(rows_with_any_data - non_empty[index], 0) / rows_with_any_data, 4)
            if rows_with_any_data
            else None,
            "cell_data_types": dict(data_types[index]),
            "distinct_count": len(distinct[index]),
            "samples": sample_values[index],
        }
        for index in range(len(headers))
    ]

    counted_headers = {
        "result indicator",
        "phễu",
        "funnel",
        "objective",
        "campaign objective",
        "objective suy luận",
        "optimization goal",
        "performance goal",
        "publisher platform",
        "platform position",
        "destination",
        "conversion location",
    }
    record["value_counts"] = {}
    for index, header in enumerate(headers, start=1):
        if norm(header) not in counted_headers:
            continue
        counts: Counter[str] = Counter()
        for row in range(header_row + 1, ws.max_row + 1):
            value_text = text(ws.cell(row, index).value)
            counts[value_text or "[BLANK]"] += 1
        record["value_counts"][header] = dict(counts.most_common())

    record["target_columns"] = [
        column
        for column in record["columns"]
        if norm(column["header"]) in TARGET_FIELDS
        or any(token in norm(column["header"]) for token in ("objective", "optimization", "publisher", "platform", "placement", "destination", "phễu", "funnel", "product", "sku", " id"))
    ]

    id_columns = []
    for index, header in enumerate(headers, start=1):
        if norm(header) in ID_FIELDS or norm(header).endswith(" id"):
            id_columns.append(inspect_id_column(ws, header_row, index, header))
    record["id_columns"] = id_columns

    spend_col = canonical_field(headers, ["Amount spent (VND)", "Amount spent"])
    result_col = canonical_field(headers, ["Results"])
    result_indicator_col = canonical_field(headers, ["Result indicator"])
    name_col = canonical_field(headers, ["Campaign name", "Ad set name", "Ad name"])
    if spend_col:
        spend_total = 0.0
        spend_total_named_rows = 0.0
        spend_total_blank_name_rows = 0.0
        spend_rows = 0
        spend_named_rows = 0
        spend_blank_name_rows = 0
        spent_without_results = 0
        spent_without_indicator = 0
        spent_samples: list[dict[str, str]] = []
        for row in range(header_row + 1, ws.max_row + 1):
            value = ws.cell(row, spend_col).value
            try:
                spend = float(value) if value not in (None, "") else 0.0
            except (TypeError, ValueError):
                continue
            if spend <= 0:
                continue
            spend_rows += 1
            spend_total += spend
            name = text(ws.cell(row, name_col).value) if name_col else ""
            if name:
                spend_named_rows += 1
                spend_total_named_rows += spend
            else:
                spend_blank_name_rows += 1
                spend_total_blank_name_rows += spend
            result = ws.cell(row, result_col).value if result_col else None
            indicator = ws.cell(row, result_indicator_col).value if result_indicator_col else None
            if result in (None, "", 0, 0.0):
                spent_without_results += 1
            if indicator in (None, ""):
                spent_without_indicator += 1
            if len(spent_samples) < 5:
                spent_samples.append(
                    {
                        "row": str(row),
                        "name": name,
                        "spend": text(value),
                        "results": text(result),
                        "result_indicator": text(indicator),
                    }
                )
        record["spend_profile"] = {
            "positive_spend_rows": spend_rows,
            "sum_including_possible_total_rows": round(spend_total, 6),
            "positive_spend_named_rows": spend_named_rows,
            "sum_named_rows": round(spend_total_named_rows, 6),
            "positive_spend_blank_name_rows": spend_blank_name_rows,
            "sum_blank_name_rows": round(spend_total_blank_name_rows, 6),
            "positive_spend_rows_with_zero_or_blank_results": spent_without_results,
            "positive_spend_rows_with_blank_result_indicator": spent_without_indicator,
            "samples": spent_samples,
        }
    return record


def inspect_workbook(path: Path, scope: str) -> dict[str, Any]:
    record: dict[str, Any] = {
        "path": str(path),
        "scope": scope,
        "exists": path.exists(),
        "month": month_from_path(path),
        "grain": grain_from_path(path),
    }
    if not path.exists():
        return record
    try:
        workbook = load_workbook(path, read_only=False, data_only=False)
        record["sheets"] = [inspect_sheet(path, ws) for ws in workbook.worksheets]
        workbook.close()
    except Exception as exc:
        record["error"] = repr(exc)
    return record


def build_summary(records: list[dict[str, Any]]) -> dict[str, Any]:
    field_locations: defaultdict[str, list[dict[str, str]]] = defaultdict(list)
    preferred_profiles: list[dict[str, Any]] = []
    for record in records:
        for sheet in record.get("sheets", []):
            for header in sheet.get("headers", []):
                normalized = norm(header)
                if any(
                    token in normalized
                    for token in (
                        "objective",
                        "optimization",
                        "performance goal",
                        "publisher",
                        "platform",
                        "placement",
                        "destination",
                        "conversion location",
                        "phễu",
                        "funnel",
                        "product",
                        "sku",
                        " id",
                    )
                ):
                    field_locations[header].append({"path": record["path"], "sheet": sheet["sheet"]})
        if "meta_ads\\preferred_candidate" in record["path"]:
            first_sheet = record.get("sheets", [{}])[0]
            preferred_profiles.append(
                {
                    "path": record["path"],
                    "month": record["month"],
                    "grain": record["grain"],
                    "header_row": first_sheet.get("header_row"),
                    "data_rows_with_any_value": first_sheet.get("data_rows_with_any_value"),
                    "headers": first_sheet.get("headers", []),
                    "spend_profile": first_sheet.get("spend_profile"),
                    "id_columns": first_sheet.get("id_columns", []),
                }
            )
    return {
        "field_locations": dict(field_locations),
        "preferred_profiles": preferred_profiles,
    }


def main() -> None:
    global WORKSPACE, RAW_ROOT, OUTPUT_ROOT, WORK_ROOT, OUTPUT_JSON, SELECTED_OUTPUTS
    parser = argparse.ArgumentParser(description="Audit Joycat sources from a portable workspace root.")
    parser.add_argument("--root", type=Path, default=WORKSPACE)
    parser.add_argument("--output", type=Path)
    args = parser.parse_args()
    WORKSPACE = args.root.expanduser().resolve()
    RAW_ROOT = WORKSPACE / "01_inputs" / "joycat" / "raw"
    OUTPUT_ROOT = WORKSPACE / "01_inputs" / "joycat"
    WORK_ROOT = WORKSPACE / "02_work" / "joycat" / "dataset_audit"
    OUTPUT_JSON = (args.output or WORK_ROOT / "source_audit_v2.json").resolve()
    SELECTED_OUTPUTS = [
        OUTPUT_ROOT / "JOYCAT_CAMPAIGN_3_THANG_CO_CHI_PHI.xlsx",
        OUTPUT_ROOT / "JOYCAT_CAMPAIGN_3_THANG_CO_CHI_PHI_OBJECTIVE_DEMO.xlsx",
        OUTPUT_ROOT / "JOYCAT_SHOPEE_PRODUCTS_2026-08-25.xlsx",
    ]
    raw_files = sorted(RAW_ROOT.rglob("*.xlsx"), key=lambda path: str(path).casefold())
    records = [inspect_workbook(path, "raw") for path in raw_files]
    records.extend(inspect_workbook(path, "output") for path in SELECTED_OUTPUTS)
    result = {
        "generated_at": datetime.now().isoformat(timespec="seconds"),
        "record_count": len(records),
        "records": records,
        "summary": build_summary(records),
    }
    WORK_ROOT.mkdir(parents=True, exist_ok=True)
    OUTPUT_JSON.write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps({"record_count": len(records), "output": str(OUTPUT_JSON)}, ensure_ascii=False))


if __name__ == "__main__":
    main()
