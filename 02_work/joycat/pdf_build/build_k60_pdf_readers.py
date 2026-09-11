from __future__ import annotations

import argparse
import csv
import os
import re
from pathlib import Path
from typing import Iterable
from xml.sax.saxutils import escape

from openpyxl import load_workbook
from PIL import Image as PILImage
from pypdf import PdfReader, PdfWriter
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_LEFT
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.pdfgen import canvas
from reportlab.platypus import (
    KeepTogether,
    PageBreak,
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)
from reportlab.lib.utils import ImageReader


NAVY = colors.HexColor("#17324D")
BLUE = colors.HexColor("#2F6FED")
TEAL = colors.HexColor("#0F8B8D")
LIGHT = colors.HexColor("#F3F6F9")
MID = colors.HexColor("#D9E2EC")
DARK = colors.HexColor("#1F2933")
MUTED = colors.HexColor("#52606D")
ORANGE = colors.HexColor("#F59E0B")
RED = colors.HexColor("#B42318")


def register_fonts() -> None:
    font_dir = Path(os.environ.get("WINDIR", r"C:\Windows")) / "Fonts"
    candidates = {
        "K60": font_dir / "arial.ttf",
        "K60-Bold": font_dir / "arialbd.ttf",
        "K60-Italic": font_dir / "ariali.ttf",
    }
    for name, path in candidates.items():
        if not path.exists():
            raise FileNotFoundError(f"Missing required font: {path}")
        pdfmetrics.registerFont(TTFont(name, str(path)))


def styles():
    base = getSampleStyleSheet()
    return {
        "title": ParagraphStyle(
            "K60Title",
            parent=base["Title"],
            fontName="K60-Bold",
            fontSize=26,
            leading=31,
            textColor=NAVY,
            alignment=TA_LEFT,
            spaceAfter=12,
        ),
        "subtitle": ParagraphStyle(
            "K60Subtitle",
            parent=base["Normal"],
            fontName="K60",
            fontSize=13,
            leading=19,
            textColor=MUTED,
            spaceAfter=10,
        ),
        "h1": ParagraphStyle(
            "K60H1",
            parent=base["Heading1"],
            fontName="K60-Bold",
            fontSize=18,
            leading=23,
            textColor=NAVY,
            spaceBefore=10,
            spaceAfter=8,
        ),
        "h2": ParagraphStyle(
            "K60H2",
            parent=base["Heading2"],
            fontName="K60-Bold",
            fontSize=14,
            leading=19,
            textColor=TEAL,
            spaceBefore=8,
            spaceAfter=6,
        ),
        "body": ParagraphStyle(
            "K60Body",
            parent=base["BodyText"],
            fontName="K60",
            fontSize=10.5,
            leading=15.5,
            textColor=DARK,
            spaceAfter=6,
        ),
        "small": ParagraphStyle(
            "K60Small",
            parent=base["BodyText"],
            fontName="K60",
            fontSize=8.5,
            leading=12,
            textColor=MUTED,
        ),
        "table": ParagraphStyle(
            "K60Table",
            parent=base["BodyText"],
            fontName="K60",
            fontSize=8.1,
            leading=11,
            textColor=DARK,
        ),
        "table_head": ParagraphStyle(
            "K60TableHead",
            parent=base["BodyText"],
            fontName="K60-Bold",
            fontSize=8.4,
            leading=11,
            textColor=colors.white,
        ),
        "center": ParagraphStyle(
            "K60Center",
            parent=base["BodyText"],
            fontName="K60",
            fontSize=10,
            leading=15,
            alignment=TA_CENTER,
            textColor=MUTED,
        ),
    }


def p(text: str, style) -> Paragraph:
    return Paragraph(text, style)


def bullets(items: Iterable[str], st) -> list:
    out = []
    for item in items:
        out.append(Paragraph(f"- {escape(item)}", st))
    return out


def note_box(text: str, st, color=TEAL):
    t = Table([[Paragraph(text, st)]], colWidths=[165 * mm])
    t.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, -1), LIGHT),
                ("BOX", (0, 0), (-1, -1), 1, color),
                ("LEFTPADDING", (0, 0), (-1, -1), 10),
                ("RIGHTPADDING", (0, 0), (-1, -1), 10),
                ("TOPPADDING", (0, 0), (-1, -1), 9),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 9),
            ]
        )
    )
    return t


def doc_footer(label: str):
    def _draw(c: canvas.Canvas, doc):
        c.saveState()
        c.setStrokeColor(MID)
        c.line(20 * mm, 13 * mm, 190 * mm, 13 * mm)
        c.setFont("K60", 7.5)
        c.setFillColor(MUTED)
        c.drawString(20 * mm, 8 * mm, label)
        c.drawRightString(190 * mm, 8 * mm, f"Trang {doc.page}")
        c.restoreState()

    return _draw


def build_doc(path: Path, story: list, label: str, pagesize=A4) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    doc = SimpleDocTemplate(
        str(path),
        pagesize=pagesize,
        leftMargin=20 * mm,
        rightMargin=20 * mm,
        topMargin=20 * mm,
        bottomMargin=19 * mm,
        title=label,
        author="Duy - Meta Agentic BA",
    )
    footer = doc_footer(label)
    doc.build(story, onFirstPage=footer, onLaterPages=footer)


def add_cover(story: list, st, title: str, subtitle: str, status: str = "K60 Reader") -> None:
    story.extend(
        [
            Spacer(1, 22 * mm),
            p(status.upper(), ParagraphStyle("eyebrow", parent=st["small"], fontName="K60-Bold", textColor=BLUE, fontSize=9.5, leading=12)),
            Spacer(1, 4 * mm),
            p(escape(title), st["title"]),
            p(escape(subtitle), st["subtitle"]),
            Spacer(1, 7 * mm),
            note_box(
                "Vai trò: tài liệu học và phương pháp. Không dùng reader này thay cho dữ liệu hoặc bằng chứng business của Joycat.",
                st["body"],
            ),
            Spacer(1, 10 * mm),
            p("Meta Agentic BA | 09/09/2026", st["small"]),
        ]
    )


def table(data: list[list], widths: list[float], st, repeat_rows=1) -> Table:
    rendered = []
    for r, row in enumerate(data):
        cell_style = st["table_head"] if r == 0 else st["table"]
        rendered.append([cell if isinstance(cell, Paragraph) else p(escape(str(cell)), cell_style) for cell in row])
    t = Table(rendered, colWidths=widths, repeatRows=repeat_rows, hAlign="LEFT")
    t.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), NAVY),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("GRID", (0, 0), (-1, -1), 0.35, MID),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, LIGHT]),
                ("LEFTPADDING", (0, 0), (-1, -1), 6),
                ("RIGHTPADDING", (0, 0), (-1, -1), 6),
                ("TOPPADDING", (0, 0), (-1, -1), 5),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
            ]
        )
    )
    return t


def build_reading_guide(output: Path, st) -> None:
    story = []
    add_cover(
        story,
        st,
        "K60 - Hướng dẫn đọc và tra cứu",
        "Bản đồ chọn tài liệu cho Duy, CEO/reviewer và AI mới nhận workspace.",
    )
    story.append(PageBreak())
    story.append(p("1. Đọc nhanh trong 5 phút", st["h1"]))
    story.extend(
        bullets(
            [
                "K60 là nguồn phương pháp; Joycat là case dùng để kiểm nghiệm phương pháp.",
                "Muốn phân tích Joycat phải đọc Context Joycat và dữ liệu Joycat sau khi đọc K60.",
                "Không xem filename, tên Campaign hoặc slide ví dụ là bằng chứng cấu hình thật.",
                "PDF reader giúp đọc thuận tiện; source gốc vẫn là nguồn truy cứu chi tiết.",
            ],
            st["body"],
        )
    )
    story.append(p("2. Chọn tài liệu theo câu hỏi", st["h1"]))
    rows = [
        ["Nếu Duy đang cần...", "Đọc file", "Kết quả mong đợi"],
        ["Biết bắt đầu phân tích từ đâu", "01 - Analytical Thinking", "Problem statement, decomposition, evidence và điểm dừng"],
        ["Hiểu KPI, data quality và maturity", "01 - Analytical Thinking", "Phân biệt operational metric, business value, quality và maturity"],
        ["Hiểu sample data/model", "02 - Data Sample Catalog", "Biết table, grain, key và bài tập nào có thể dùng"],
        ["Học AI Fluency session", "03 - AI Fluency Status", "Biết source hiện thiếu gì và cần xin lại gì"],
        ["Xem sample Power BI", "04 - Power BI Catalog", "Biết file kỹ thuật đang có và điều kiện mở/export"],
        ["Chọn chart, màu và kể chuyện", "05 - Visualization", "Dùng chart có mục đích và trình bày insight rõ"],
    ]
    story.append(table(rows, [46 * mm, 48 * mm, 71 * mm], st))
    story.append(Spacer(1, 5 * mm))
    story.append(p("3. Chuỗi áp dụng vào Joycat", st["h1"]))
    story.append(
        note_box(
            "Yêu cầu của Duy -> Context hiện hành -> câu hỏi phân tích -> K60 để chọn phương pháp -> source Joycat để kiểm nghiệm -> Power BI để lưu bằng chứng -> kết luận có limitation.",
            st["body"],
            BLUE,
        )
    )
    story.append(p("4. Bốn nhãn bắt buộc khi ghi nhận kiến thức", st["h1"]))
    rows = [
        ["Nhãn", "Ý nghĩa", "Ví dụ"],
        ["Đã xác minh từ nguồn", "Tài liệu hoặc dữ liệu trực tiếp hỗ trợ", "Slide K60 nói về data-quality dimensions"],
        ["Owner đã xác nhận", "Người có thẩm quyền xác nhận", "Cậu Sinh xác nhận mục đích bài tập"],
        ["Suy luận", "Diễn giải cần kiểm nghiệm", "Nhóm Campaign có thể thuộc một phễu"],
        ["To be updated", "Thiếu source hoặc đang mâu thuẫn", "Chưa có MP4/transcript AI Fluency"],
    ]
    story.append(table(rows, [35 * mm, 62 * mm, 68 * mm], st))
    story.append(p("5. Những việc reader không làm thay", st["h1"]))
    story.extend(
        bullets(
            [
                "Không thay Data Dictionary của Joycat.",
                "Không xác nhận Campaign objective, funnel hoặc customer journey Joycat.",
                "Không tính Business ROAS khi chưa có GMV business cùng scope và period.",
                "Không xác nhận ETL readiness nếu chưa kiểm grain, key, field và lineage.",
                "Không tự tạo recommendation từ một metric đứng riêng lẻ.",
            ],
            st["body"],
        )
    )
    build_doc(output, story, "K60 Reading Guide")


def build_analytical_front(path: Path, st) -> None:
    story = []
    add_cover(
        story,
        st,
        "K60 - Analytical Thinking và Logic Tree",
        "Reader tổng hợp PowerPoint về Data & Analytics metrics với tài liệu Problem Solving 101.",
    )
    story.append(PageBreak())
    story.append(p("Cách đọc", st["h1"]))
    story.extend(
        bullets(
            [
                "Phần A: 30 slide về operational metrics, business value, data quality, data literacy và data-management maturity.",
                "Phần B: tài liệu Problem Solving 101, giữ nguyên 68 trang nguồn.",
                "Khi áp dụng: bắt đầu từ câu hỏi/decision, phân rã logic, kiểm source rồi mới chọn metric và visualization.",
            ],
            st["body"],
        )
    )
    story.append(p("Điểm cần nhớ khi dùng cho Joycat", st["h1"]))
    rows = [
        ["Dùng được", "Phải xác minh lại"],
        ["Cách phân rã vấn đề và kiểm tra giả thuyết", "KPI target và benchmark của Joycat"],
        ["Khung operational/business/data-quality metrics", "Field, grain, key và attribution của dataset Joycat"],
        ["Cách dùng scorecard/heat map", "Ngưỡng đỏ/vàng/xanh và owner của quyết định"],
    ]
    story.append(table(rows, [82 * mm, 83 * mm], st))
    story.append(PageBreak())
    story.append(p("Mục lục reader", st["h1"]))
    story.append(p("Phần A - Data & Analytics Operational, Data Quality, and Data Management Maturity Metrics", st["body"]))
    story.append(p("Phần B - Problem Solving 101", st["body"]))
    story.append(Spacer(1, 6 * mm))
    story.append(note_box("Nội dung hai phần sau được giữ nguyên từ source. Trang reader tổng hợp không phải bản dịch hoặc diễn giải thay nguồn.", st["body"], ORANGE))
    build_doc(path, story, "K60 Analytical Thinking")


def build_divider(path: Path, st, title: str, subtitle: str) -> None:
    story = []
    add_cover(story, st, title, subtitle, status="Phần tiếp theo")
    build_doc(path, story, title)


def merge_analytical(output: Path, front: Path, ppt_pdf: Path, divider: Path, source_pdf: Path) -> None:
    writer = PdfWriter()
    page_starts = []
    for label, src in [
        ("Hướng dẫn sử dụng", front),
        ("Phần A - Data & Analytics Metrics", ppt_pdf),
        ("Phần B - Problem Solving 101", divider),
        ("Problem Solving 101 - nội dung nguồn", source_pdf),
    ]:
        page_starts.append((label, len(writer.pages)))
        reader = PdfReader(str(src))
        for page in reader.pages:
            writer.add_page(page)
    for label, page_no in page_starts:
        writer.add_outline_item(label, page_no)
    writer.add_metadata(
        {
            "/Title": "K60 - Analytical Thinking and Logic Tree",
            "/Author": "Duy - Meta Agentic BA",
            "/Subject": "Derived reader from K60 source materials",
        }
    )
    with output.open("wb") as f:
        writer.write(f)


def first_nonempty_row(ws, max_scan: int = 12) -> list[str]:
    for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row or 1, max_scan), values_only=True):
        values = ["" if v is None else str(v) for v in row]
        if any(v.strip() for v in values):
            return values
    return []


def unique_csv_files(root: Path) -> list[Path]:
    seen = set()
    out = []
    for pth in root.rglob("*"):
        if pth.is_file() and pth.suffix.lower() == ".csv":
            key = str(pth.resolve()).lower()
            if key not in seen:
                seen.add(key)
                out.append(pth)
    return sorted(out)


def csv_header(path: Path) -> list[str]:
    with path.open("r", encoding="utf-8-sig", errors="replace", newline="") as f:
        sample = f.read(65536)
        f.seek(0)
        try:
            dialect = csv.Sniffer().sniff(sample, delimiters=",;|\t")
        except csv.Error:
            dialect = csv.excel
        return next(csv.reader(f, dialect), [])


def build_data_catalog(output: Path, source_root: Path, st) -> None:
    story = []
    add_cover(
        story,
        st,
        "K60 - Data Sample Catalog",
        "Bản đọc schema sơ bộ cho các Excel, CSV và Access sample; không sao chép toàn bộ dữ liệu.",
    )
    story.append(PageBreak())
    story.append(p("1. Cách dùng catalogue", st["h1"]))
    story.extend(
        bullets(
            [
                "Dùng để chọn sample phù hợp khi học grain, key, star schema và data model.",
                "Rows/columns chỉ là quan sát kỹ thuật, chưa xác nhận business definition.",
                "Không mang table hoặc field của Contoso/QlikView sang Joycat nếu chưa mapping lại.",
            ],
            st["body"],
        )
    )
    story.append(p("2. Excel workbooks", st["h1"]))
    for path in sorted(source_root.rglob("*.xlsx")):
        rel = path.relative_to(source_root)
        rows = [["Sheet", "Rows x Columns", "Dòng đầu có dữ liệu"]]
        try:
            wb = load_workbook(path, read_only=True, data_only=True)
            for ws in wb.worksheets:
                hdr = first_nonempty_row(ws)
                hdr_text = ", ".join(v for v in hdr[:12] if v)
                if len(hdr) > 12:
                    hdr_text += ", ..."
                rows.append([ws.title, f"{ws.max_row:,} x {ws.max_column:,}", hdr_text or "Không quan sát được"])
            wb.close()
        except Exception as exc:
            rows.append(["Không đọc được", "-", str(exc)])
        story.append(
            KeepTogether(
                [
                    p(escape(str(rel)), st["h2"]),
                    table(rows, [35 * mm, 32 * mm, 98 * mm], st),
                    Spacer(1, 4 * mm),
                ]
            )
        )

    story.append(p("3. CSV files", st["h1"]))
    csv_rows = [["File", "Dung lượng", "Header quan sát được"]]
    for path in unique_csv_files(source_root):
        try:
            headers = csv_header(path)
            header_text = ", ".join(headers[:14]) + (", ..." if len(headers) > 14 else "")
        except Exception as exc:
            header_text = f"Không đọc được: {exc}"
        csv_rows.append([str(path.relative_to(source_root)), f"{path.stat().st_size / 1024:,.1f} KB", header_text])
    story.append(table(csv_rows, [61 * mm, 25 * mm, 79 * mm], st))
    story.append(Spacer(1, 5 * mm))
    story.append(p("4. Access database", st["h1"]))
    access_files = list(source_root.rglob("*.accdb"))
    if access_files:
        for path in access_files:
            story.append(p(f"<b>{escape(path.name)}</b> - {path.stat().st_size / (1024 * 1024):,.1f} MB", st["body"]))
            story.append(note_box("Chưa audit table/schema trong phase tạo PDF. Muốn dùng phải mở bằng Access/driver phù hợp và lập data dictionary riêng.", st["body"], ORANGE))
    rows = [
        ["Khái niệm", "Sample phù hợp", "Câu hỏi tự kiểm"],
        ["Fact và dimension", "Contoso / invoice detail / Styles", "Mỗi dòng đại diện cho điều gì? Key nào nối được?"],
        ["Target vs actual", "Sales Target + sales fact", "Hai bảng có cùng grain và period không?"],
        ["Promotion", "Sales Promotion", "Discount ở grain Date x ProductCategory có phân bổ được không?"],
        ["Inventory", "On Hand History", "Snapshot tồn kho có được cộng qua ngày không?"],
    ]
    story.append(
        KeepTogether(
            [
                p("5. Gợi ý bài học", st["h1"]),
                table(rows, [35 * mm, 54 * mm, 76 * mm], st),
            ]
        )
    )
    build_doc(output, story, "K60 Data Sample Catalog")


def build_ai_fluency_status(output: Path, source_root: Path, st) -> None:
    candidates = list(source_root.rglob("*.drivedownload"))
    story = []
    add_cover(
        story,
        st,
        "K60 - AI Fluency Session Status",
        "Bản ghi mức sẵn sàng của Session E-learning Agentic AI by Anthropic.",
    )
    story.append(PageBreak())
    story.append(p("Kết luận hiện tại", st["h1"]))
    story.append(note_box("Chưa có video hoặc transcript đọc được. File hiện tại có đuôi .mp4.drivedownload nên không được coi là nội dung session đã sẵn sàng.", st["body"], RED))
    story.append(p("Source quan sát được", st["h1"]))
    rows = [["File", "Dung lượng", "Trạng thái"]]
    for path in candidates:
        rows.append([path.name, f"{path.stat().st_size / (1024 * 1024):,.1f} MB", "Placeholder tải xuống; chưa phải MP4"])
    if len(rows) == 1:
        rows.append(["Không tìm thấy", "-", "Chưa có source"])
    story.append(table(rows, [78 * mm, 30 * mm, 57 * mm], st))
    story.append(p("Cần bổ sung", st["h1"]))
    story.extend(
        bullets(
            [
                "Tải lại file MP4 hoàn chỉnh hoặc xin transcript/subtitle.",
                "Ghi rõ source, ngày/version và quyền sử dụng.",
                "Sau khi có source thật mới tạo PDF note hoặc transcript reader.",
            ],
            st["body"],
        )
    )
    story.append(p("Phần vẫn làm được", st["h1"]))
    story.append(p("Các reader Analytical Thinking, Data Sample và Visualization vẫn dùng độc lập. Việc thiếu video không chặn việc đọc các source khác.", st["body"]))
    build_doc(output, story, "K60 AI Fluency Session Status")


def build_powerbi_catalog(output: Path, source_root: Path, st) -> None:
    pbix_files = list(source_root.rglob("*.pbix"))
    story = []
    add_cover(
        story,
        st,
        "K60 - Power BI Sample Catalog",
        "Bản hướng dẫn truy cứu sample Python scatter plot trong Power BI.",
    )
    story.append(PageBreak())
    story.append(p("Source kỹ thuật", st["h1"]))
    rows = [["File", "Dung lượng", "Trạng thái"]]
    for path in pbix_files:
        rows.append([str(path.relative_to(source_root)), f"{path.stat().st_size / (1024 * 1024):,.1f} MB", "PBIX; chưa export report pages"])
    if len(rows) == 1:
        rows.append(["Không tìm thấy PBIX", "-", "Chưa có source"])
    story.append(table(rows, [93 * mm, 27 * mm, 45 * mm], st))
    story.append(Spacer(1, 5 * mm))
    story.append(note_box("Filename gợi ý đây là sample tạo scatter plot bằng Python, nhưng filename không đủ để xác nhận data model, measure, visual hoặc code thật bên trong.", st["body"], ORANGE))
    story.append(p("Quy trình đọc đúng", st["h1"]))
    story.extend(
        bullets(
            [
                "Mở bản sao PBIX bằng Power BI Desktop; không sửa source gốc.",
                "Chụp danh sách report pages, tables, relationships, measures và Python script.",
                "Export report pages thành PDF nếu source mở ổn định.",
                "Ghi lại Power BI version, warning, external dependency và refresh requirement.",
                "Chỉ sau đó mới cập nhật reader này thành bản nội dung đầy đủ.",
            ],
            st["body"],
        )
    )
    story.append(p("Khi áp dụng vào Joycat", st["h1"]))
    rows = [
        ["Học từ sample", "Không được sao chép mù"],
        ["Cách tổ chức visual và analytical question", "Measure, table, relationship và threshold"],
        ["Cách dùng scatter plot để so sánh hai numeric metrics", "Kết luận tương quan hoặc nguyên nhân"],
        ["Cách lưu evidence trong report", "Business grain và attribution của Joycat"],
    ]
    story.append(table(rows, [82 * mm, 83 * mm], st))
    build_doc(output, story, "K60 Power BI Sample Catalog")


def slide_number(path: Path) -> int:
    match = re.search(r"(\d+)", path.stem)
    return int(match.group(1)) if match else 10**9


def build_visualization(output: Path, source_root: Path) -> None:
    images = sorted(source_root.joinpath("session 16-18 visualization").glob("*.jpg"), key=slide_number)
    page_w, page_h = landscape(A4)
    c = canvas.Canvas(str(output), pagesize=(page_w, page_h), pageCompression=1)
    c.setTitle("K60 - Session 16-18 Visualization")
    c.setAuthor("Duy - Meta Agentic BA")

    c.setFillColor(NAVY)
    c.rect(0, 0, page_w, page_h, fill=1, stroke=0)
    c.setFillColor(colors.white)
    c.setFont("K60-Bold", 26)
    c.drawString(22 * mm, page_h - 42 * mm, "K60 - Session 16-18 Visualization")
    c.setFont("K60", 13)
    c.setFillColor(colors.HexColor("#D9E8FF"))
    c.drawString(22 * mm, page_h - 55 * mm, f"Reader gồm {len(images)} ảnh slide, giữ thứ tự số file nguồn.")
    c.setFont("K60", 10)
    c.setFillColor(colors.white)
    text = c.beginText(22 * mm, page_h - 78 * mm)
    text.setLeading(16)
    for line in [
        "Dùng để học: chart choice, màu sắc, dashboard, bias và storytelling.",
        "Không dùng slide học tập làm bằng chứng cho kết luận Joycat.",
        "Các số slide nguồn không liên tục; đây là đặc điểm source, không phải lỗi ghép PDF.",
    ]:
        text.textLine(f"- {line}")
    c.drawText(text)
    c.setFont("K60", 8)
    c.setFillColor(colors.HexColor("#D9E8FF"))
    c.drawRightString(page_w - 18 * mm, 12 * mm, "Trang 1")
    c.showPage()

    for index, img_path in enumerate(images, start=2):
        c.setFillColor(colors.white)
        c.rect(0, 0, page_w, page_h, fill=1, stroke=0)
        header_h = 12 * mm
        footer_h = 10 * mm
        c.setFillColor(NAVY)
        c.rect(0, page_h - header_h, page_w, header_h, fill=1, stroke=0)
        c.setFillColor(colors.white)
        c.setFont("K60-Bold", 8.5)
        c.drawString(12 * mm, page_h - 8 * mm, f"K60 Visualization | {img_path.name}")

        with PILImage.open(img_path) as im:
            iw, ih = im.size
        max_w = page_w - 18 * mm
        max_h = page_h - header_h - footer_h - 8 * mm
        scale = min(max_w / iw, max_h / ih)
        dw, dh = iw * scale, ih * scale
        x = (page_w - dw) / 2
        y = footer_h + (max_h - dh) / 2
        c.drawImage(ImageReader(str(img_path)), x, y, width=dw, height=dh, preserveAspectRatio=True, mask="auto")

        c.setFillColor(MUTED)
        c.setFont("K60", 7.5)
        c.drawString(12 * mm, 5 * mm, "Source: session 16-18 visualization")
        c.drawRightString(page_w - 12 * mm, 5 * mm, f"Trang {index}")
        c.showPage()
    c.save()


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--source-root", type=Path, required=True)
    parser.add_argument("--output-dir", type=Path, required=True)
    parser.add_argument("--tmp-dir", type=Path, required=True)
    parser.add_argument("--ppt-pdf", type=Path, required=True)
    args = parser.parse_args()

    register_fonts()
    st = styles()
    args.output_dir.mkdir(parents=True, exist_ok=True)
    args.tmp_dir.mkdir(parents=True, exist_ok=True)

    build_reading_guide(args.output_dir / "00_K60_READING_GUIDE.pdf", st)

    front = args.tmp_dir / "analytical_front.pdf"
    divider = args.tmp_dir / "problem_solving_divider.pdf"
    build_analytical_front(front, st)
    build_divider(divider, st, "Phần B - Problem Solving 101", "68 trang source gốc được ghép nguyên trạng sau trang phân cách này.")
    source_pdf = next(args.source_root.rglob("problem_solving_101.pdf"))
    merge_analytical(
        args.output_dir / "01_K60_ANALYTICAL_THINKING_AND_LOGIC_TREE.pdf",
        front,
        args.ppt_pdf,
        divider,
        source_pdf,
    )

    build_data_catalog(args.output_dir / "02_K60_DATA_SAMPLE_CATALOG.pdf", args.source_root, st)
    build_ai_fluency_status(args.output_dir / "03_K60_AI_FLUENCY_SESSION_STATUS.pdf", args.source_root, st)
    build_powerbi_catalog(args.output_dir / "04_K60_POWER_BI_SAMPLE_CATALOG.pdf", args.source_root, st)
    build_visualization(args.output_dir / "05_K60_VISUALIZATION_SESSION_16_18.pdf", args.source_root)

    expected = [
        "00_K60_READING_GUIDE.pdf",
        "01_K60_ANALYTICAL_THINKING_AND_LOGIC_TREE.pdf",
        "02_K60_DATA_SAMPLE_CATALOG.pdf",
        "03_K60_AI_FLUENCY_SESSION_STATUS.pdf",
        "04_K60_POWER_BI_SAMPLE_CATALOG.pdf",
        "05_K60_VISUALIZATION_SESSION_16_18.pdf",
    ]
    missing = [name for name in expected if not (args.output_dir / name).exists()]
    if missing:
        raise RuntimeError(f"Missing outputs: {missing}")
    print("CREATED")
    for name in expected:
        path = args.output_dir / name
        print(f"{name}|{path.stat().st_size}")


if __name__ == "__main__":
    main()
